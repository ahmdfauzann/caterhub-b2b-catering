import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_documentation():
    doc = Document()

    # Page Setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles Setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x25, 0x2A)

    # Custom Heading Colors
    BRAND_COLOR = RGBColor(0xEA, 0x58, 0x0C) # Brand Orange
    NAVY_COLOR = RGBColor(0x0F, 0x17, 0x2A)  # Navy 900
    GRAY_COLOR = RGBColor(0x47, 0x55, 0x69)  # Slate 600

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = BRAND_COLOR
        p.paragraph_format.space_after = Pt(4)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.color.rgb = GRAY_COLOR
        p.paragraph_format.space_after = Pt(20)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = NAVY_COLOR
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = BRAND_COLOR
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)

    def add_h3(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = NAVY_COLOR
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = NAVY_COLOR
        r_text = p.add_run(text)
        if italic:
            r_text.font.italic = True

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = NAVY_COLOR
        p.add_run(text)

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)

        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        cell._tc.get_or_add_tcPr().append(shading)

        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="EA580C"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Document Header
    add_title("DOKUMENTASI KODE, ALUR SISTEM & PANDUAN LIVE CODING")
    add_subtitle("CaterHub - B2B Catering Marketplace Platform (Laravel 11 Pure Handcrafted)")

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nama Proyek", "CaterHub (B2B Catering Marketplace)"),
        ("Teknologi Framework", "Laravel 11 (Pure Code, No Auto-CRUD Builders)"),
        ("Developer / Pengembang", "Ahmad Fauzan"),
        ("Tujuan Proyek", "Tes Kemampuan Bidang Programming - PT Jasamedika Transmedic")
    ]
    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        c1.paragraphs[0].add_run(k).font.bold = True
        c2.paragraphs[0].add_run(v)
        c1._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFF7ED"/>'))
        c2._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>'))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # BAB 1
    add_h1("BAB 1. PENDAHULUAN & ARSITEKTUR BISNIS B2B")
    add_body("Aplikasi CaterHub adalah platform B2B Catering Marketplace yang dirancang khusus untuk menjembatani kebutuhan makan siang harian, event kantor, dan meeting direksi antara perusahaan (Customer) dengan vendor katering profesional (Merchant).")
    
    add_h2("1.1 Fitur Utama & Penyederhanaan Spesifikasi")
    add_bullet("Menggunakan 2 Role Utama: merchant (Mitra Vendor Katering) dan customer (Klien Perusahaan Kantor). Membuang role superadmin agar aplikasi fokus, bersih, dan sesuai dengan skenario pengerjaan tes.", "Sistem Akses 2 Role: ")
    add_bullet("Dibuat menggunakan murni native Laravel 11 tanpa dependency library auto-CRUD seperti Filament, Backpack, Nova, atau Voyager. Seluruh Controller, Model, Migration, dan Blade Template ditulis secara manual (handcrafted).", "Standar Handcrafted Code: ")
    add_bullet("Menggunakan metode Corporate Billing (Invoice B2B pasca-pengiriman) yang diterbitkan secara otomatis dengan tanggal jatuh tempo (due date).", "Sistem Pembayaran B2B: ")
    add_bullet("Sesi waktu pengiriman disesuaikan menjadi rentang jam presisi (contoh: 08:00 - 08:30 WIB, 11:30 - 12:00 WIB, 12:00 - 12:30 WIB, 15:30 - 16:00 WIB, dll).", "Jadwal Pengiriman Jam-ke-Jam: ")
    add_bullet("Menyediakan perintah artisan kustom `php artisan db:export-sql` yang menghasilkan berkas `database.sql` di root project secara otomatis dari database SQLite/MySQL.", "Penyediaan Export Database: ")

    # BAB 2
    add_h1("BAB 2. PENJELASAN ALUR APLIKASI SECARA DETAIL (APPLICATION FLOW)")
    add_body("Berikut adalah alur perjalanan sistem (user journey & business workflow) dalam aplikasi CaterHub dari pendaftaran akun hingga pelunasan invoice:")

    add_h2("2.1 Alur Registrasi & Autentikasi User")
    add_bullet("Pengguna membuka menu registrasi di `/register`. Sistem menampilkan pilihan jenis akun: Klien Kantor (Customer) atau Mitra Katering (Merchant).", "1. Pemilihan Peran: ")
    add_bullet("Formulir pendaftaran kantor di `/register/customer` meminta data Nama PIC, Email, Password, Nama Perusahaan (`company_name`), Nomor Telepon, Alamat Kantor, dan Jumlah Karyawan.", "2. Registrasi Customer: ")
    add_bullet("Formulir pendaftaran vendor di `/register/merchant` meminta data Owner, Email, Password, Nama Katering, Jenis Masakan (`cuisine_type`), Alamat Dapur, Kota, dan Min. Order.", "3. Registrasi Merchant: ")
    add_bullet("Sistem mengautentikasi email & password di AuthController. Jika role `merchant`, di-redirect ke `/merchant/dashboard`. Jika role `customer`, di-redirect ke `/customer/dashboard`.", "4. Login & Role Redirection: ")

    add_h2("2.2 Alur Eksplorasi & Pencarian Katering")
    add_bullet("Halaman Beranda (`/`) menampilkan Hero Banner, Form Pencarian Cepat, Kategori Makanan Popular, Mitra Favorit, dan Menu Terlaris.", "1. Beranda Utama: ")
    add_bullet("Halaman `/search` memproses query string seperti `q` (nama menu/katering), `city` (kota), `category` (slug), dan `max_price` (harga max).", "2. Filter & Pencarian: ")
    add_bullet("Di halaman `/catering/{slug}`, customer dapat melihat profil katering, alamat dapur, jenis masakan, rating ulasan, serta daftar menu per kategori.", "3. Detail Merchant & Menu: ")

    add_h2("2.3 Alur Keranjang Belanja & Checkout B2B")
    add_bullet("Customer memilih menu dan porsi (minimal porsi divalidasi). Data disimpan dalam `session('cart')`.", "1. Penambahan ke Keranjang: ")
    add_bullet("Di halaman `/customer/checkout`, customer mengisi Tanggal Pengiriman, Sesi Waktu Pengiriman (contoh: `11:30 - 12:00 WIB`), Alamat Pengiriman, Catatan, dan metode `Corporate Billing`.", "2. Formulir Checkout: ")
    add_bullet("CustomerPortalController menghitung Subtotal, Pajak PPN 10%, Delivery Fee, dan Grand Total. Record transaksi disimpan di tabel `orders`, `order_items`, dan `invoices`.", "3. Pembuatan Pesanan & Invoice: ")

    add_h2("2.4 Alur Pengelolaan Pesanan oleh Vendor")
    add_bullet("Vendor login dan masuk ke `/merchant/dashboard` untuk monitoring pendapatan dan pesanan.", "1. Monitoring Pesanan: ")
    add_bullet("Vendor mengelola menu (CRUD) di `/merchant/menus` termasuk upload foto.", "2. Manajemen Menu: ")
    add_bullet("Vendor update status bertahap melalui `PATCH /merchant/orders/{id}/status`: Pending -> Confirmed -> Preparing -> Delivering -> Delivered -> Completed.", "3. Update Status Operasional: ")

    add_h2("2.5 Alur Penerbitan Invoice & Pembayaran")
    add_bullet("Invoice B2B diterbitkan otomatis dengan nomor `INV-YYYYMMDD-CAT-XXXX` di rute `/invoice/{invoiceNumber}`.", "1. Penomoran Invoice Otomatis: ")
    add_bullet("Rute `/invoice/{invoiceNumber}/print` menyediakan tampilan web invoice printable.", "2. Cetak Web Invoice: ")

    add_h2("2.6 Alur Ulasan & Rating")
    add_bullet("Customer memberikan rating 1-5 bintang dan ulasan komentar setelah pesanan `delivered`/`completed`.", "1. Ulasan Customer: ")
    add_bullet("Nilai rata-rata `rating_avg` pada `merchant_profiles` dihitung ulang secara otomatis.", "2. Kalkulasi Rating Otomatis: ")

    # BAB 3
    add_h1("BAB 3. STRUKTUR DIREKTORI & KELOMPOK FILE")
    add_bullet("Pengendali logika bisnis (AuthController, CustomerPortalController, MerchantPortalController, HomeController, InvoiceController).", "1. app/Http/Controllers/: ")
    add_bullet("Representasi tabel database (User, MerchantProfile, CustomerProfile, Category, Menu, Order, OrderItem, Invoice, Review).", "2. app/Models/: ")
    add_bullet("Pengaman hak akses role (`RoleMiddleware.php`).", "3. app/Http/Middleware/: ")
    add_bullet("File migrasi skema tabel DDL & seeder data (`DatabaseSeeder.php`).", "4. database/: ")
    add_bullet("Tampilan UI Blade Template + Tailwind CSS.", "5. resources/views/: ")
    add_bullet("File ekspor database ANSI SQL standar.", "6. database.sql: ")

    # BAB 4
    add_h1("BAB 4. BEDAH BARIS KODE FILE PENENTU")

    add_h2("4.1 Controller Utama")
    add_h3("A. AuthController.php")
    add_code_block(
"// 1. Registrasi Customer\n"
"$user = User::create(['name' => $request->name, 'email' => $request->email, 'password' => Hash::make($request->password), 'role' => 'customer']);\n"
"CustomerProfile::create(['user_id' => $user->id, 'company_name' => $request->company_name, 'office_address' => $request->office_address, 'city' => $request->city]);\n\n"
"// 2. Redireksi Role\n"
"return $user->isMerchant() ? redirect()->route('merchant.dashboard') : redirect()->route('customer.dashboard');"
    )

    add_h3("B. CustomerPortalController.php")
    add_code_block(
"// Pembuatan Order & Invoice di Checkout\n"
"$order = Order::create([\n"
"    'order_code' => 'ORD-' . date('Ymd') . '-' . sprintf('%03d', rand(1, 999)),\n"
"    'customer_id' => auth()->id(), 'merchant_id' => $merchantId,\n"
"    'delivery_date' => $request->delivery_date, 'delivery_time' => $request->delivery_time,\n"
"    'total_amount' => $subtotal, 'tax_amount' => $tax, 'delivery_fee' => $deliveryFee, 'grand_total' => $grandTotal,\n"
"    'payment_method' => 'Corporate Billing', 'status' => 'pending', 'payment_status' => 'unpaid',\n"
"]);\n\n"
"Invoice::create(['invoice_number' => 'INV-' . date('Ymd') . '-CAT-' . sprintf('%04d', $order->id), 'order_id' => $order->id, 'issue_date' => now(), 'due_date' => now()->addDays(7), 'amount' => $grandTotal, 'status' => 'unpaid']);"
    )

    add_h3("C. MerchantPortalController.php")
    add_code_block(
"// Update Status Pesanan oleh Vendor\n"
"public function updateOrderStatus(Request $request, $id) {\n"
"    $order = Order::where('merchant_id', $this->getMerchantProfile()->id)->findOrFail($id);\n"
"    $order->update(['status' => $request->status]);\n"
"    if (in_array($request->status, ['delivered', 'completed']) && $order->payment_status === 'unpaid') {\n"
"        $order->update(['payment_status' => 'paid']);\n"
"        if ($order->invoice) { $order->invoice->update(['status' => 'paid']); }\n"
"    }\n"
"    return redirect()->back()->with('success', 'Status pesanan berhasil diperbarui.');\n"
"}"
    )

    # BAB 5
    add_h1("BAB 5. KESIMPULAN & PETUNJUK PENGUJIAN")
    add_body("Seluruh pengujian unit & feature test berjalan 100% PASSING (4 passed, 7 assertions). Gunakan akun customer@jasamedika.com (password) atau berkah@catering.com (password) untuk pengujian lokal.")

    # NEW SECTION: BAB 6
    add_h1("BAB 6. PANDUAN PENGEMBANGAN FITUR (ADD, REMOVE, SEARCH & MODIFY)")
    add_body("Bab ini disusun khusus untuk membimbing Anda jika penguji interview meminta Anda menambah fitur baru, menghapus bagian tertentu, atau mencari data spesifik secara langsung saat live coding.")

    add_h2("6.1 CARA MENAMBAHKAN FITUR BARU (HOW TO ADD A FEATURE)")

    add_h3("Skenario A: Menambahkan Fitur Kode Kupon Diskon Promo pada Checkout")
    add_body("Jika penguji meminta Anda menambah kolom Kode Diskon/Kupon saat checkout:")
    add_bullet("1. Buat Migration Tabel Kupon (atau tambahkan kolom discount_amount pada tabel orders):", "Langkah 1: Migration -> ")
    add_code_block("php artisan make:migration add_discount_to_orders_table --table=orders")
    add_body("Di dalam migration file, tambahkan:")
    add_code_block("$table->decimal('discount_amount', 12, 2)->default(0)->after('delivery_fee');")
    add_bullet("2. Update Model Order.php dengan memasukkan 'discount_amount' ke dalam array $fillable.", "Langkah 2: Model -> ")
    add_bullet("3. Update CustomerPortalController.php pada method processCheckout():", "Langkah 3: Controller -> ")
    add_code_block(
"$discount = 0;\n"
"if ($request->filled('coupon_code') && $request->coupon_code === 'PROMO10') {\n"
"    $discount = $subtotal * 0.10; // Diskon 10%\n"
"}\n"
"$grandTotal = ($subtotal - $discount) + $tax + $deliveryFee;"
    )
    add_bullet("4. Update Blade View resources/views/customer/checkout.blade.php dengan menambahkan input text kode promo.", "Langkah 4: View -> ")

    add_h3("Skenario B: Menambahkan Upload Foto Profil / Avatar User")
    add_body("Jika penguji meminta menambah fitur foto profil avatar pada user:")
    add_bullet("1. Tambahkan kolom avatar pada migration users atau run `php artisan make:migration add_avatar_to_users_table`.", "Langkah 1: ")
    add_bullet("2. Di CustomerPortalController.php / AuthController.php, tambahkan proses upload file:", "Langkah 2: ")
    add_code_block(
"if ($request->hasFile('avatar')) {\n"
"    $path = $request->file('avatar')->store('avatars', 'public');\n"
"    $user->update(['avatar' => $path]);\n"
"}"
    )

    add_h2("6.2 CARA MENGHAPUS ATAU MEMATIKAN FITUR (HOW TO REMOVE A FEATURE)")

    add_h3("Skenario A: Menghapus Fitur Ulasan / Rating Katering")
    add_body("Jika penguji meminta untuk mematikan atau menghapus fitur Review & Rating:")
    add_bullet("Hapus blok rute `Route::post('/orders/{id}/review', ...)` di `routes/web.php`.", "1. Rute: ")
    add_bullet("Di `CustomerPortalController.php`, hapus method `storeReview()`.", "2. Controller: ")
    add_bullet("Di view `resources/views/customer/orders/show.blade.php`, hapus modal atau form bintang rating.", "3. View: ")

    add_h3("Skenario B: Mematikan Biaya Pengiriman (Free Ongkir)")
    add_body("Jika penguji meminta membuat biaya pengiriman Rp 0 (Free Shipping):")
    add_body("Di `CustomerPortalController.php` method `processCheckout()`, ubah variabel `$deliveryFee = 0;`.")

    add_h2("6.3 CARA MENCARI DAN MEMFILTER DATA SPESIFIK (HOW TO QUERY & SEARCH DATA)")
    add_body("Berikut adalah contoh-contoh query Eloquent yang sering diminta oleh penguji live coding:")

    add_h3("1. Mencari Pesanan Berdasarkan Rentang Tanggal (Date Range Filter)")
    add_code_block(
"// Filter pesanan dari tanggal A sampai tanggal B\n"
"$orders = Order::whereBetween('delivery_date', ['2026-10-01', '2026-10-31'])\n"
"               ->where('merchant_id', $merchantId)\n"
"               ->get();"
    )

    add_h3("2. Mencari Pesanan yang Belum Lunas (Unpaid Invoices)")
    add_code_block(
"$unpaidOrders = Order::with('invoice')\n"
"    ->where('payment_status', 'unpaid')\n"
"    ->latest()\n"
"    ->get();"
    )

    add_h3("3. Menghitung Total Pendapatan Vendor Bulan Ini")
    add_code_block(
"$totalRevenue = Order::where('merchant_id', $merchantId)\n"
"    ->where('status', 'completed')\n"
"    ->whereMonth('created_at', now()->month)\n"
"    ->sum('grand_total');"
    )

    add_h3("4. Mencari Menu Makanan Terlaris (Top Selling Items)")
    add_code_block(
"$topMenus = Menu::where('merchant_profile_id', $merchantId)\n"
"    ->orderBy('sales_count', 'desc')\n"
"    ->take(5)\n"
"    ->get();"
    )

    # NEW SECTION: BAB 7
    add_h1("BAB 7. PREDIKSI PERTANYAAN & STRATEGI INTERVIEW LIVE CODING")
    add_body("Bab ini berisi bocoran pertanyaan yang paling sering ditanyakan oleh Penguji/Tech Lead saat live coding tes beserta jawaban & strategi menjawabnya:")

    add_h2("7.1 Prediksi Pertanyaan Penguji & Jawaban Terbaik")

    add_h3("Q1: 'Mengapa Anda menggunakan 2 Role dan tidak menggunakan Superadmin?'")
    add_bullet("Sesuai instruksi soal dan prinsip YAGNI (You Aren't Gonna Need It), aplikasi difokuskan pada alur utama transaksi B2B antara Kantor (Customer) dan Vendor (Merchant). Menghilangkan Superadmin membuat kodebase jauh lebih bersih, ringan, dan mudah di-maintain tanpa mengurangi pemenuhan syarat soal.", "Jawaban Terbaik: ")

    add_h3("Q2: 'Bagaimana cara aplikasi mencegah User mengakses halaman yang bukan haknya?'")
    add_bullet("Aplikasi menggunakan custom middleware `RoleMiddleware.php` yang dipasang pada grup rute di `web.php`. Middleware akan mengecek apakah `auth()->user()->role` sesuai dengan parameter role yang diizinkan (misal `role:merchant` atau `role:customer`). Jika mencoba bypass URL, user otomatis di-redirect ke beranda dengan pesan alert.", "Jawaban Terbaik: ")

    add_h3("Q3: 'Bagaimana jika aplikasi ini ingin dipindahkan dari SQLite ke MySQL di server produksi?'")
    add_bullet("Sangat mudah! Karena Laravel menggunakan ORM Eloquent dan DB Migration Abstraction, kita hanya perlu mengubah konfigurasi file `.env` dari `DB_CONNECTION=sqlite` menjadi `DB_CONNECTION=mysql` serta mengisikan `DB_DATABASE`, `DB_USERNAME`, dan `DB_PASSWORD`. Setelah itu jalankan `php artisan migrate:fresh --seed` atau import file `database.sql` yang telah kita siapkan.", "Jawaban Terbaik: ")

    add_h3("Q4: 'Bagaimana alur kerja Keranjang Belanja (Cart) di aplikasi ini?'")
    add_bullet("Keranjang belanja menggunakan `Session` (`session()->get('cart')`). Saat customer menekan 'Tambah ke Keranjang', item disimpan ke dalam array session di server. Keunggulannya: cepat, tidak membebani database dengan data draf temporary, dan otomatis dibersihkan setelah transaksi checkout berhasil dibuat (`session()->forget('cart')`).", "Jawaban Terbaik: ")

    add_h3("Q5: 'Bagaimana cara memastikan porsi minimal yang dipesan sesuai dengan aturan vendor?'")
    add_bullet("Validasi dilakukan di 2 lapisan (Defense in Depth). Di frontend (Blade template), input number diset `min=\"{{ $menu->min_portion }}\"`. Di backend (`CustomerPortalController.php`), kita melakukan pengecekan `$request->quantity >= $menu->min_portion`. Jika kurang, controller mengembalikan pesan error validasi.", "Jawaban Terbaik: ")

    add_h2("7.2 Tips & Mentalitas Saat Live Coding")
    add_bullet("Selalu jelaskan apa yang sedang Anda ketik. Contoh: 'Sekarang saya akan menambahkan method updateStatus di MerchantPortalController untuk mengubah status order...'", "1. Berpikir Keras Secara Lisan (Think Out Loud): ")
    add_bullet("Gunakan terminal `php artisan tinker` untuk mengetes query Eloquent dengan cepat sebelum memasukkannya ke dalam kode Controller.", "2. Manfaatkan Artisan Tinker: ")
    add_bullet("Jika terjadi error saat live coding, jangan panik! Gulung layar terminal ke bagian atas atau cek file `storage/logs/laravel.log` untuk membaca pesan error utama. Penguji sangat menyukai kandidat yang tenang dan mahir membaca error log.", "3. Mahir Membaca Error Trace: ")

    # Save document
    file_path = "c:\\xampp\\htdocs\\apkcatering\\Dokumentasi_Aplikasi_CaterHub.docx"
    doc.save(file_path)
    print(f"Document updated successfully at: {file_path}")

if __name__ == "__main__":
    create_documentation()
